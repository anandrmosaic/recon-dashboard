"""
ShipBob mails, lite version: pure Python (no pandas/numpy) -- same two mails, same Excel layout, same email text
as mailer.py, built from engine_lite's `Ship` objects instead of a pandas DataFrame.

Used on Render (where pandas broke production once -- see engine_lite.py). Manual drafting only: nothing here
runs on a timer or a watcher. A draft is only ever created when the user clicks the button in the page
(POST /api/mailer/<kind>/draft), matching the local app's manual-drafting mode.

State lives in the same files the full mailer uses (state/mailer_state.json, state/mailer_log.csv), so whichever
engine you're looking at (local or Render) sees the same "already drafted" / "already sent" history -- as long as
both point at the same `shipbob_live/state/` folder. On Render that folder is empty at first and fills in as you
use it there; it is NOT synced from your PC (state/ is git-ignored and local-only by design).

  build(kind, ships, meta)   Excel + email text
  create_draft(result)       saves a Gmail draft with the Excel attached (updates the same-day draft if one exists)
  mark_sent(kind)            you confirm you sent it (draft -> sent in the log)

There is no auto_run() here on purpose -- Render never drafts anything by itself.
"""
import base64
import csv
import html as _html
import json
import os
from collections import Counter
from datetime import date, datetime
from email.mime.application import MIMEApplication
from email.mime.multipart import MIMEMultipart
from email.mime.text import MIMEText

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from . import engine_lite as engine

OUT_DIR = os.path.join(engine.BASE, "output", "mailer")
STATE_FILE = os.path.join(engine.BASE, "state", "mailer_state.json")
MAIL_LOG = os.path.join(engine.BASE, "state", "mailer_log.csv")
GMAIL_SCOPES = ["https://www.googleapis.com/auth/gmail.compose"]

KINDS = ("potential", "claims")
MAIL_KINDS = ("potential", "pending", "raised")
TITLES = {"potential": "Potential in-transit loss", "claims": "Claims follow-up"}
FILES = {"potential": "Potential_InTransit_Loss", "claims": "Claims_Followup"}

BASE_COLS = [("Shipment ID", 13), ("Store Order ID", 18), ("Import Date", 13), ("Age", 12), ("Carrier", 16),
             ("Tracking Information", 28), ("Delivery Status", 42)]
EXTRA_COLS = {"potential": [],
              "pending": [("Expected Claim $", 14), ("Claims remark", 40)],
              "raised": [("Claim Age", 16), ("SLA status", 20), ("Expected Claim $", 14), ("Claims remark", 40)]}
HEADERS = {"Age": "Days Since Import", "Tracking Information": "Tracking #", "Expected Claim $": "Expected Claim ($)",
           "Mailed": "First mailed", "Claim Age": "Days Since Claim Raised"}
BANDS = [("16-30 days", 16, 30), ("31-60 days", 31, 60), ("61-90 days", 61, 90), ("91+ days", 91, 99999)]


# ------------------------------------------------------------------ state (same files as mailer.py)
def load_state():
    st = {}
    if os.path.exists(STATE_FILE):
        with open(STATE_FILE, encoding="utf-8") as f:
            st = json.load(f)
    if "kinds" not in st:
        st = {"kinds": {}}
    return st


def kstate(kind):
    return load_state()["kinds"].get(kind, {})


def save_kstate(kind, **changes):
    st = load_state()
    st["kinds"].setdefault(kind, {}).update(changes)
    os.makedirs(os.path.dirname(STATE_FILE), exist_ok=True)
    with open(STATE_FILE, "w", encoding="utf-8") as f:
        json.dump(st, f, indent=1)
    return st["kinds"][kind]


def status(cfg=None):
    cfg = cfg or engine.load_config()
    out = {"gmail_connected": gmail_connected(cfg), "kinds": {}}
    for k in KINDS:
        s = dict(kstate(k))
        s.pop("ids_by_kind", None)
        s["excel_exists"] = bool(s.get("excel_path")) and os.path.exists(s["excel_path"])
        s["mail_to"] = _to(k, cfg)
        s["title"] = TITLES[k]
        out["kinds"][k] = s
    return out


def _to(kind, cfg):
    return cfg.get(f"mail_to_{kind}") or cfg.get("mail_to", "")


def _cc(kind, cfg):
    return cfg.get(f"mail_cc_{kind}") or cfg.get("mail_cc", "")


# ------------------------------------------------------------------ mail log (pandas-free; same CSV as mailer.py)
def _read_log_rows():
    if not os.path.exists(MAIL_LOG):
        return []
    with open(MAIL_LOG, newline="", encoding="utf-8") as f:
        rows = list(csv.DictReader(f))
    for r in rows:
        r.setdefault("Status", "drafted")
        r.setdefault("Kind", "potential")
    return rows


def load_mail_log():
    """{kind: {shipment id: {"date": ..., "status": ...}}}. A shipment ever marked sent counts as sent.
    Same dedup rule as mailer.py: prefer status=sent, then earliest Sent On, one entry per (Kind, Shipment ID)."""
    rows = _read_log_rows()
    out = {k: {} for k in MAIL_KINDS}
    best = {}
    for r in rows:
        key = (r["Kind"], r["Shipment ID"])
        rank = 1 if r["Status"] == "sent" else 0
        cur = best.get(key)
        if cur is None or (rank, ) > (cur[0], ) or (rank == cur[0] and r["Sent On"] < cur[1]["Sent On"]):
            best[key] = (rank, r)
    for (kind, sid), (_rank, r) in best.items():
        try:
            out.setdefault(kind, {})[int(float(sid))] = {"date": r["Sent On"], "status": r["Status"]}
        except (ValueError, TypeError):
            continue
    return out


def append_mail_log(shipment_ids, sent_on, status="drafted", kind="potential"):
    os.makedirs(os.path.dirname(MAIL_LOG), exist_ok=True)
    existing = _read_log_rows()
    seen = {(r["Shipment ID"], r["Sent On"], r["Status"], r["Kind"]) for r in existing}
    new_rows = []
    for sid in shipment_ids:
        row = {"Shipment ID": str(sid), "Sent On": sent_on, "Status": status, "Kind": kind}
        key = (row["Shipment ID"], row["Sent On"], row["Status"], row["Kind"])
        if key not in seen:
            seen.add(key)
            new_rows.append(row)
    all_rows = existing + new_rows
    with open(MAIL_LOG, "w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=["Shipment ID", "Sent On", "Status", "Kind"])
        w.writeheader()
        w.writerows(all_rows)


def _mailed_text(entry):
    if not entry:
        return ""
    d = datetime.strptime(str(entry["date"])[:10], "%Y-%m-%d").strftime("%d %b")
    return f'{"Sent" if entry["status"] == "sent" else "Drafted"} {d}'


# ------------------------------------------------------------------ build
def sla_status(claim_age, cfg):
    lo, hi = cfg.get("claim_sla_days", [45, 60])
    return "Over %d days" % hi if claim_age > hi else ("%d-%d days" % (lo, hi) if claim_age >= lo else "Within %d days" % lo)


def parts(ships, kind, cfg):
    """The order lists inside a mail: {part kind: [Ship, ...]}, oldest (highest Age) first."""
    by_age = lambda sel: sorted(sel, key=lambda s: (s.age is None, -(s.age or 0)))
    if kind == "potential":
        return {"potential": by_age([s for s in ships if s.potential])}
    pend = by_age([s for s in ships if s.sub == "Pending to claim"])
    raised = by_age([s for s in ships if s.sub == "Claim raised but not received"])
    return {"pending": pend, "raised": raised}


def _mailed_lookup(part):
    log = load_mail_log()
    return log.get(part, {})


def _cell_val(s, col, mailed_map, cfg):
    if col == "Shipment ID":
        return s.id
    if col == "Store Order ID":
        return s.store
    if col == "Import Date":
        return s.import_dt
    if col == "Age":
        return s.age
    if col == "Carrier":
        return s.carrier
    if col == "Tracking Information":
        return s.tracking
    if col == "Delivery Status":
        return s.delivery
    if col == "Expected Claim $":
        return s.exp_claim
    if col == "Claims remark":
        return s.remark
    if col == "Claim Age":
        return s.claim_age
    if col == "SLA status":
        return sla_status(s.claim_age or 0, cfg)
    if col == "Mailed":
        return _mailed_text(mailed_map.get(s.id))
    return None


def _sheet(wb, title, ships, part, mailed_map, cfg, first=False):
    ws = wb.active if first else wb.create_sheet()
    ws.title = title[:31]
    cols = BASE_COLS + EXTRA_COLS[part] + [("Mailed", 16)]
    head_fill = PatternFill("solid", fgColor="1F3864")
    for i, (c, w) in enumerate(cols, 1):
        cell = ws.cell(row=1, column=i, value=HEADERS.get(c, c))
        cell.font = Font(name="Arial", bold=True, color="FFFFFF")
        cell.fill = head_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(i)].width = w
    for r, s in enumerate(ships, 2):
        for i, (c, _w) in enumerate(cols, 1):
            v = _cell_val(s, c, mailed_map, cfg)
            cell = ws.cell(row=r, column=i, value=v)
            cell.font = Font(name="Arial", size=10)
            if c == "Import Date":
                cell.number_format = "dd-mmm-yyyy"
            elif c == "Expected Claim $":
                cell.number_format = "#,##0.00"
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{max(len(ships) + 1, 2)}"
    ws.row_dimensions[1].height = 30
    return ws


SHEET_NAMES = {"potential": "Potential loss", "pending": "Pending to claim", "raised": "Raised, not received"}


def count_new(kind, pp, data_date):
    ks = kstate(kind)
    same_day = ks.get("drafted_at", "")[:10] if ks.get("drafted_for") == data_date else None
    n = 0
    for part, ships in pp.items():
        mailed_map = _mailed_lookup(part)
        for s in ships:
            e = mailed_map.get(s.id)
            if e is None or (same_day and e["status"] == "drafted" and str(e["date"])[:10] == same_day):
                n += 1
    return n


def build(kind, ships, meta, cfg=None):
    """Build the Excel and the email text for one mail. Does not touch Gmail or the log."""
    cfg = cfg or engine.load_config()
    pp = parts(ships, kind, cfg)
    n_all = sum(len(d) for d in pp.values())
    n_new = count_new(kind, pp, f"{meta['as_of']:%Y-%m-%d}")
    os.makedirs(OUT_DIR, exist_ok=True)
    path = os.path.join(OUT_DIR, f"ShipBob_{FILES[kind]}_{meta['as_of']:%Y-%m-%d}.xlsx")
    wb = Workbook()
    for i, (part, d) in enumerate(pp.items()):
        mailed_map = _mailed_lookup(part)
        _sheet(wb, f"{SHEET_NAMES[part]} ({len(d)})", d, part, mailed_map, cfg, first=(i == 0))
    wb.save(path)
    subject, blocks = (_email_potential if kind == "potential" else _email_claims)(pp, n_new, meta, cfg)
    body, html = render_plain(blocks), render_html(blocks)
    with open(path.replace(".xlsx", "_email.txt"), "w", encoding="utf-8") as f:
        f.write(f"Subject: {subject}\n\n{body}\n")
    ids = {part: [s.id for s in d] for part, d in pp.items()}
    value = round(sum(s.exp_claim or 0.0 for p, d in pp.items() if p != "potential" for s in d), 2) if kind == "claims" else None
    lo = cfg.get("claim_sla_days", [45, 60])[0]
    aged = sum(1 for s in pp["raised"] if (s.claim_age or 0) >= lo) if kind == "claims" else None
    save_kstate(kind, excel_path=path, excel_built_at=datetime.now().isoformat(timespec="seconds"),
                data_date=f"{meta['as_of']:%Y-%m-%d}", new_count=n_new, all_count=n_all, value=value, aged=aged,
                pending_count=len(pp.get("pending", [])) if kind == "claims" else None,
                raised_count=len(pp.get("raised", [])) if kind == "claims" else None,
                subject=subject, body=body, html=html, ids_by_kind=ids)
    return {"kind": kind, "path": path, "all": n_all, "new": n_new, "subject": subject, "body": body, "html": html,
            "ids_by_kind": ids}


# ------------------------------------------------------------------ email text (identical to mailer.py)
def render_plain(blocks):
    out = []
    for blk in blocks:
        kind, val = blk[0], blk[1]
        if kind in ("P", "H", "L"):
            out.append(val)
        elif kind == "UL":
            out.append("\n".join(f"  - {i}" for i in val))
        else:
            out.append("\n".join(f"  {chr(97 + n)}. {i}" for n, i in enumerate(val)))
    text = ""
    for n, blk in enumerate(blocks):
        text += out[n]
        if n < len(blocks) - 1:
            tight = blk[0] in ("H", "L") or (blk[0] == "P" and len(blk) > 2 and blk[2])
            text += "\n" if tight else "\n\n"
    return text


def render_html(blocks):
    h = lambda t: _html.escape(str(t)).replace("\n", "<br>")
    parts_ = []
    for blk in blocks:
        kind, val = blk[0], blk[1]
        tight = len(blk) > 2 and blk[2]
        if kind == "P":
            parts_.append(f'<p style="margin:0 0 {4 if tight else 12}px">{h(val)}</p>')
        elif kind in ("H", "L"):
            parts_.append(f'<p style="margin:{16 if kind == "H" else 12}px 0 4px"><b>{h(val)}</b></p>')
        elif kind == "UL":
            parts_.append('<ul style="margin:0 0 12px;padding-left:22px">' + "".join(f"<li>{h(i)}</li>" for i in val) + "</ul>")
        else:
            parts_.append('<ol type="a" style="margin:0 0 12px;padding-left:22px">' + "".join(f"<li>{h(i)}</li>" for i in val) + "</ol>")
    return ('<div style="font-family:Arial,Helvetica,sans-serif;font-size:14px;line-height:1.5;color:#222">'
            + "".join(parts_) + "</div>")


def _counts(items):
    """items: list of raw values (e.g. [s.stage for s in allf]) -- same order/format as pandas' value_counts()."""
    c = Counter(items)
    return [f"{k}: {v}" for k, v in c.most_common()]


def _band_counts(ships):
    return [f"{b}: {sum(1 for s in ships if lo <= (s.age or -1) <= hi)}" for b, lo, hi in BANDS]


def _newtxt(n_all, n_new):
    if n_new == n_all:
        return ""
    if n_new == 0:
        return " All of them were shared earlier and are still open."
    return f" {n_new} of them are new since our last mail; the rest were shared earlier and are still open."


def _sign(cfg):
    return ("P", f"Thanks,\n{cfg.get('mail_signature', '')}")


def _email_potential(pp, n_new, meta, cfg):
    allf = pp["potential"]
    A, T = cfg["alert_from_day"], cfg["transit_days"]
    new_in_subject = f" ({n_new} new)" if 0 < n_new < len(allf) else ""
    subject = (f"ShipBob D2C - {len(allf)} orders not delivered {A}-{T} days after import"
               f"{new_in_subject} - {meta['as_of']:%d %b %Y}")
    return subject, [
        ("P", "Hi Team,"),
        ("P", f"Please find attached the D2C orders that are still not delivered {A}-{T} days after import "
              f"(data as of {meta['as_of']:%d %b %Y}).{_newtxt(len(allf), n_new)} They move into dispute after {T} days, "
              f"so we would appreciate a status check with the carriers before then."),
        ("L", "By latest carrier status:"), ("UL", _counts([s.stage for s in allf])),
        ("L", "By carrier:"), ("UL", _counts([s.carrier or "Not assigned" for s in allf])),
        ("P", "Could you please share an update / ETA for each order, and flag any that are lost so we can proceed with the claim?"),
        _sign(cfg),
    ]


def _email_claims(pp, n_new, meta, cfg):
    pend, raised = pp["pending"], pp["raised"]
    T = cfg["transit_days"]
    lo, hi = cfg.get("claim_sla_days", [45, 60])
    n_all = len(pend) + len(raised)
    subject = (f"ShipBob D2C - claims follow-up: {len(pend)} pending to claim, {len(raised)} raised not received "
               f"({meta['as_of']:%d %b %Y})")
    blocks = [("P", "Hi Team,"),
              ("P", f"Two updates on our D2C claims (data as of {meta['as_of']:%d %b %Y}). "
                    f"The order-wise lists are in the attached sheet, one tab for each.{_newtxt(n_all, n_new)}")]
    sec = 1
    if len(pend):
        blocks += [
            ("H", f"{sec}. Pending to claim: unable to raise on the portal"),
            ("P", f"We are unable to raise claims on the ShipBob portal for {len(pend)} orders. These orders are not "
                  f"delivered and are more than {T} days past import."),
            ("L", "Ageing (days since order import):"), ("UL", _band_counts(pend)),
            ("P", "Could you please:", True),
            ("OL", ["Let us know the reason we are unable to raise the claim on these orders, and",
                    "Share the current status of each order, or raise the claim from your side if that is possible."])]
        sec += 1
    if len(raised):
        within = sum(1 for s in raised if (s.claim_age or 0) < lo)
        window = sum(1 for s in raised if lo <= (s.claim_age or 0) <= hi)
        over = sum(1 for s in raised if (s.claim_age or 0) > hi)
        blocks += [
            ("H", f"{sec}. Claim raised, amount not received"),
            ("P", f"We have raised claims for {len(raised)} orders and the amount has not been received yet."),
            ("L", "Ageing since the claim was raised:"),
            ("UL", [f"Under {lo} days: {within}", f"{lo}-{hi} days: {window}", f"Over {hi} days: {over}"])]
        if window + over:
            blocks.append(("P", f"As per the claim SLA of {lo}-{hi} days, the timeline has passed for {window + over} of these "
                                f"orders and the amount is still not received ({over} of them are already past {hi} days). "
                                f"Could you please resolve these on priority and share the status / expected credit date?"))
        if within:
            blocks.append(("P", f"For the remaining {within} orders (still within the SLA), could you please check "
                                f"whether there is any update on them?"))
    blocks.append(_sign(cfg))
    return subject, blocks


# ------------------------------------------------------------------ gmail
# Render has no local config.json / writable token file and no browser for interactive consent, so there a
# pre-authorized token is read from the GMAIL_TOKEN_JSON env var instead (same pattern as auth.py). Locally,
# shipbob_live/config.json's gmail_token_file + the interactive browser flow is used, same as mailer.py.
def _env_creds():
    token_json = os.environ.get("GMAIL_TOKEN_JSON")
    if not token_json:
        return None
    from google.auth.transport.requests import Request
    from google.oauth2.credentials import Credentials
    creds = Credentials.from_authorized_user_info(json.loads(token_json), GMAIL_SCOPES)
    if creds.expired and creds.refresh_token:
        creds.refresh(Request())
    return creds


def gmail_connected(cfg=None):
    creds = _env_creds()
    if creds is not None:
        return bool(creds.valid or creds.refresh_token)
    cfg = cfg or engine.load_config()
    token = cfg.get("gmail_token_file")
    if not token or not os.path.exists(token):
        return False
    try:
        from google.oauth2.credentials import Credentials
        creds = Credentials.from_authorized_user_file(token, GMAIL_SCOPES)
        return bool(creds and (creds.valid or creds.refresh_token))
    except Exception:
        return False


def _gmail_creds(cfg, interactive=False):
    creds = _env_creds()
    if creds is not None:
        return creds

    from google.auth.transport.requests import Request
    from google.oauth2.credentials import Credentials
    from google_auth_oauthlib.flow import InstalledAppFlow

    token = cfg.get("gmail_token_file")
    if not token:
        raise RuntimeError("Gmail is not set up here. On Render, set the GMAIL_TOKEN_JSON environment variable; "
                            "locally, set gmail_token_file in shipbob_live/config.json.")
    creds = Credentials.from_authorized_user_file(token, GMAIL_SCOPES) if os.path.exists(token) else None
    if creds and not creds.valid and creds.refresh_token:
        creds.refresh(Request())
    if not creds or not creds.valid:
        if not interactive:
            raise RuntimeError("Gmail is not connected. Click Connect Gmail on the dashboard once.")
        flow = InstalledAppFlow.from_client_secrets_file(cfg["credentials_file"], GMAIL_SCOPES)
        creds = flow.run_local_server(port=0)
    with open(token, "w") as f:
        f.write(creds.to_json())
    return creds


def connect_gmail(cfg=None):
    if os.environ.get("GMAIL_TOKEN_JSON"):
        return bool(_env_creds())
    _gmail_creds(cfg or engine.load_config(), interactive=True)
    return True


def create_draft(result, cfg=None):
    """Save the mail as a Gmail draft (not sent). Same same-day-update behaviour as mailer.py."""
    from googleapiclient.discovery import build as gbuild

    cfg = cfg or engine.load_config()
    kind = result["kind"]
    service = gbuild("gmail", "v1", credentials=_gmail_creds(cfg))
    ks = kstate(kind)
    existing = ks.get("draft_id") if ks.get("drafted_for") == ks.get("data_date") else None

    to, cc, bcc = _to(kind, cfg), _cc(kind, cfg), ""
    if existing:
        try:
            hdrs = service.users().drafts().get(userId="me", id=existing, format="metadata").execute()["message"]["payload"]["headers"]
            got = {h["name"].lower(): h["value"] for h in hdrs}
            to, cc, bcc = to or got.get("to", ""), cc or got.get("cc", ""), got.get("bcc", "")
        except Exception:
            existing = None

    msg = MIMEMultipart("mixed")
    msg["Subject"] = result["subject"]
    for name, val in (("To", to), ("Cc", cc), ("Bcc", bcc)):
        if val:
            msg[name] = val
    alt = MIMEMultipart("alternative")
    alt.attach(MIMEText(result["body"], "plain", "utf-8"))
    alt.attach(MIMEText(result["html"], "html", "utf-8"))
    msg.attach(alt)
    with open(result["path"], "rb") as f:
        part = MIMEApplication(f.read(), Name=os.path.basename(result["path"]))
    part["Content-Disposition"] = f'attachment; filename="{os.path.basename(result["path"])}"'
    msg.attach(part)
    body = {"message": {"raw": base64.urlsafe_b64encode(msg.as_bytes()).decode()}}

    draft = None
    if existing:
        try:
            draft = service.users().drafts().update(userId="me", id=existing, body=body).execute()
        except Exception:
            draft = None
    if draft is None:
        draft = service.users().drafts().create(userId="me", body=body).execute()
    for part_kind, ids in result["ids_by_kind"].items():
        append_mail_log(ids, date.today().isoformat(), "drafted", part_kind)
    save_kstate(kind, draft_id=draft["id"], drafted_at=datetime.now().isoformat(timespec="seconds"),
                drafted_for=ks.get("data_date"), drafted_count=result["all"])
    return {"draft_id": draft["id"], "url": "https://mail.google.com/mail/u/0/#drafts"}


def mark_sent(kind):
    ids_by_kind = kstate(kind).get("ids_by_kind") or {}
    n = 0
    for part_kind, ids in ids_by_kind.items():
        append_mail_log(ids, date.today().isoformat(), "sent", part_kind)
        n += len(ids)
    save_kstate(kind, sent_at=datetime.now().isoformat(timespec="seconds"))
    return n


# No auto_run() here -- Render never drafts anything on its own. Manual button only.
